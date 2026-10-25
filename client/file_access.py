from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Any


DEFAULT_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".jsonl",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".cs",
    ".java",
    ".cpp",
    ".c",
    ".h",
    ".hpp",
    ".html",
    ".css",
    ".scss",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".bat",
    ".ps1",
    ".sh",
    ".sql",
    ".log",
    ".pdf",
    ".docx",
}

SKIPPED_DIRS = {
    ".git",
    ".hg",
    ".svn",
    ".idea",
    ".vscode",
    "__pycache__",
    ".pytest_cache",
    ".mypy_cache",
    ".ruff_cache",
    "node_modules",
    ".venv",
    "venv",
    "env",
    "dist",
    "build",
    "_internal",
}

SECRET_PATTERNS = [
    re.compile(r"sk-[A-Za-z0-9_\-]{16,}"),
    re.compile(r"AIza[0-9A-Za-z_\-]{20,}"),
    re.compile(r"(?i)(api[_-]?key|secret|token|password)\s*[:=]\s*[\"']?[^\"'\s]{8,}"),
]


class FileAccessService:
    def __init__(self, base_dir: Path):
        self.base_dir = Path(base_dir).resolve()
        self.config_path = self.base_dir / "config.json"

    def handle_request(self, payload: dict[str, Any]) -> dict[str, Any]:
        operation = str(payload.get("operation") or "").strip()
        args = payload.get("args") if isinstance(payload.get("args"), dict) else {}

        try:
            if operation == "list_authorized_dirs":
                return self.list_authorized_dirs()
            if operation == "list_user_files":
                return self.list_user_files(**args)
            if operation == "search_user_files":
                return self.search_user_files(**args)
            if operation == "read_user_file":
                return self.read_user_file(**args)
            return {"error": f"Unknown file operation: {operation}"}
        except Exception as exc:
            return {"error": str(exc)}

    def list_authorized_dirs(self) -> dict[str, Any]:
        roots = self.authorized_roots()
        return {
            "authorized_dirs": [str(root) for root in roots],
            "note": (
                "Only files under these directories can be listed or read. "
                "Configure client/config.json file_access_dirs to change this list."
            ),
        }

    def list_user_files(
        self,
        directory: str = "",
        depth: int = 2,
        limit: int = 100,
        extensions: list[str] | None = None,
    ) -> dict[str, Any]:
        depth = max(0, min(int(depth), 8))
        limit = max(1, min(int(limit), 500))
        allowed_exts = self._normalize_extensions(extensions)
        roots = [self._resolve_authorized_path(directory)] if directory else self.authorized_roots()

        files = []
        scanned = 0
        for root in roots:
            for path in self._walk_files(root, depth=depth, max_files=limit - len(files)):
                scanned += 1
                if allowed_exts and path.suffix.lower() not in allowed_exts:
                    continue
                files.append(self._file_meta(path))
                if len(files) >= limit:
                    break
            if len(files) >= limit:
                break

        return {
            "files": files,
            "count": len(files),
            "scanned": scanned,
            "truncated": len(files) >= limit,
        }

    def search_user_files(
        self,
        query: str,
        directory: str = "",
        extensions: list[str] | None = None,
        limit: int = 20,
        max_files_scanned: int = 5000,
    ) -> dict[str, Any]:
        query = str(query or "").strip()
        if not query:
            return {"error": "query is required"}

        limit = max(1, min(int(limit), 100))
        max_files_scanned = max(100, min(int(max_files_scanned), 20000))
        allowed_exts = self._normalize_extensions(extensions)
        terms = [term for term in re.split(r"\s+", query.lower()) if term]
        roots = [self._resolve_authorized_path(directory)] if directory else self.authorized_roots()

        matches = []
        scanned = 0
        for root in roots:
            for path in self._walk_files(root, depth=12, max_files=max_files_scanned - scanned):
                scanned += 1
                if scanned > max_files_scanned:
                    break
                if allowed_exts and path.suffix.lower() not in allowed_exts:
                    continue
                score, snippet = self._score_path(path, terms)
                if score > 0:
                    item = self._file_meta(path)
                    item["score"] = score
                    if snippet:
                        item["snippet"] = snippet
                    matches.append(item)
                if len(matches) >= limit * 3:
                    matches.sort(key=lambda item: item["score"], reverse=True)
                    matches = matches[:limit]
            if scanned >= max_files_scanned:
                break

        matches.sort(key=lambda item: item["score"], reverse=True)
        return {
            "query": query,
            "matches": matches[:limit],
            "count": min(len(matches), limit),
            "scanned": scanned,
            "truncated": scanned >= max_files_scanned,
        }

    def read_user_file(
        self,
        path: str,
        start: int = 0,
        max_chars: int = 12000,
    ) -> dict[str, Any]:
        file_path = self._resolve_authorized_path(path)
        if not file_path.is_file():
            return {"error": f"Not a file: {file_path}"}
        if file_path.suffix.lower() not in DEFAULT_TEXT_EXTENSIONS:
            return {
                "error": f"Unsupported file extension: {file_path.suffix}",
                "hint": "Supported formats currently include text files, PDF, and .docx Word documents.",
            }

        size = file_path.stat().st_size
        if size > 2 * 1024 * 1024:
            return {
                "error": "File is too large to read directly.",
                "path": str(file_path),
                "size_bytes": size,
                "hint": "Use search_user_files or ask for a smaller text/PDF/Word file.",
            }

        start = max(0, int(start))
        max_chars = max(1000, min(int(max_chars), 50000))
        text, encoding = self._read_file_text(file_path)
        text = self._redact_secrets(text)
        start = min(start, len(text))
        end = min(len(text), start + max_chars)

        return {
            "path": str(file_path),
            "name": file_path.name,
            "extension": file_path.suffix.lower(),
            "size_bytes": size,
            "encoding": encoding,
            "start": start,
            "end": end,
            "total_chars": len(text),
            "truncated": end < len(text),
            "content": text[start:end],
        }

    def authorized_roots(self) -> list[Path]:
        roots: list[Path] = []
        roots.extend(self._roots_from_config())
        env_dirs = os.getenv("AI_PET_FILE_DIRS", "")
        if env_dirs:
            roots.extend(Path(part) for part in env_dirs.split(os.pathsep) if part.strip())

        if not roots:
            roots.append(self.base_dir.parent)

        normalized = []
        seen = set()
        for root in roots:
            try:
                resolved = Path(root).expanduser().resolve()
            except OSError:
                continue
            if not resolved.is_dir():
                continue
            key = os.path.normcase(str(resolved))
            if key in seen:
                continue
            seen.add(key)
            normalized.append(resolved)
        return normalized

    def _roots_from_config(self) -> list[Path]:
        if not self.config_path.exists():
            return []
        try:
            with open(self.config_path, "r", encoding="utf-8") as f:
                config = json.load(f)
        except (OSError, ValueError):
            return []
        dirs = config.get("file_access_dirs") or config.get("authorized_file_dirs") or []
        if isinstance(dirs, str):
            dirs = [dirs]
        if not isinstance(dirs, list):
            return []
        roots = []
        for item in dirs:
            value = str(item).strip()
            if not value:
                continue
            path = Path(value)
            roots.append(path if path.is_absolute() else self.base_dir / path)
        return roots

    def _resolve_authorized_path(self, raw_path: str) -> Path:
        raw_path = str(raw_path or "").strip()
        if not raw_path:
            raise ValueError("path is required")

        candidates = []
        path = Path(raw_path).expanduser()
        if path.is_absolute():
            candidates.append(path)
        else:
            candidates.extend(root / path for root in self.authorized_roots())

        for candidate in candidates:
            try:
                resolved = candidate.resolve()
            except OSError:
                continue
            if any(self._is_inside(resolved, root) for root in self.authorized_roots()):
                return resolved

        raise PermissionError(f"Path is outside authorized directories: {raw_path}")

    def _is_inside(self, path: Path, root: Path) -> bool:
        try:
            path_key = os.path.normcase(str(path))
            root_key = os.path.normcase(str(root))
            return os.path.commonpath([path_key, root_key]) == root_key
        except ValueError:
            return False

    def _walk_files(self, root: Path, depth: int, max_files: int):
        if max_files <= 0:
            return
        root = root.resolve()
        count = 0
        stack = [(root, 0)]
        while stack and count < max_files:
            current, current_depth = stack.pop()
            try:
                entries = sorted(current.iterdir(), key=lambda item: (not item.is_dir(), item.name.lower()))
            except OSError:
                continue
            for entry in entries:
                if count >= max_files:
                    break
                try:
                    if entry.is_dir():
                        if current_depth < depth and entry.name not in SKIPPED_DIRS:
                            stack.append((entry, current_depth + 1))
                    elif entry.is_file():
                        count += 1
                        yield entry
                except OSError:
                    continue

    def _score_path(self, path: Path, terms: list[str]) -> tuple[int, str]:
        haystack = f"{path.name} {path.suffix} {path}".lower()
        score = sum(8 for term in terms if term in haystack)
        snippet = ""

        try:
            size = path.stat().st_size
        except OSError:
            size = 0

        if path.suffix.lower() in DEFAULT_TEXT_EXTENSIONS and size <= 512 * 1024:
            try:
                text, _ = self._read_file_text(path, max_chars=4000)
            except Exception:
                text = ""
            lower = text.lower()
            for term in terms:
                index = lower.find(term)
                if index >= 0:
                    score += 3
                    start = max(0, index - 80)
                    end = min(len(text), index + 180)
                    snippet = " ".join(text[start:end].split())
                    break
        return score, snippet

    def _file_meta(self, path: Path) -> dict[str, Any]:
        stat = path.stat()
        return {
            "path": str(path),
            "name": path.name,
            "extension": path.suffix.lower(),
            "size_bytes": stat.st_size,
            "modified": int(stat.st_mtime),
        }

    def _read_text(self, path: Path) -> tuple[str, str]:
        raw = path.read_bytes()
        for encoding in ("utf-8", "utf-8-sig", "gb18030", "cp936"):
            try:
                return raw.decode(encoding), encoding
            except UnicodeDecodeError:
                continue
        return raw.decode("latin-1", errors="replace"), "latin-1"

    def _read_file_text(self, path: Path, max_chars: int | None = None) -> tuple[str, str]:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._read_pdf_text(path, max_chars=max_chars)
        if suffix == ".docx":
            return self._read_docx_text(path, max_chars=max_chars)

        text, encoding = self._read_text(path)
        if max_chars is not None:
            text = text[:max_chars]
        return text, encoding

    def _read_pdf_text(self, path: Path, max_chars: int | None = None) -> tuple[str, str]:
        try:
            from pypdf import PdfReader
        except Exception as exc:
            raise RuntimeError(f"PDF support unavailable: {exc}") from exc

        try:
            reader = PdfReader(str(path))
        except Exception as exc:
            raise RuntimeError(f"Failed to open PDF: {exc}") from exc

        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception as exc:
                raise RuntimeError(f"Encrypted PDF files are not supported: {exc}") from exc

        chunks: list[str] = []
        total = 0
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""

            text = text.strip()
            if not text:
                continue

            block = f"[Page {index}]\n{text}"
            if max_chars is not None:
                remaining = max_chars - total
                if remaining <= 0:
                    break
                block = block[:remaining]
            if block.strip():
                chunks.append(block)
                total += len(block)
            if max_chars is not None and total >= max_chars:
                break

        return "\n\n".join(chunks).strip(), "pdf"

    def _read_docx_text(self, path: Path, max_chars: int | None = None) -> tuple[str, str]:
        try:
            from docx import Document
        except Exception as exc:
            raise RuntimeError(f"Word document support unavailable: {exc}") from exc

        try:
            document = Document(str(path))
        except Exception as exc:
            raise RuntimeError(f"Failed to open Word document: {exc}") from exc

        chunks: list[str] = []
        total = 0

        def add_block(text: str) -> bool:
            nonlocal total
            text = text.strip()
            if not text:
                return False
            if max_chars is not None:
                remaining = max_chars - total
                if remaining <= 0:
                    return True
                text = text[:remaining]
            if text.strip():
                chunks.append(text)
                total += len(text)
            return max_chars is not None and total >= max_chars

        for paragraph in document.paragraphs:
            if add_block(paragraph.text or ""):
                break

        if max_chars is None or total < max_chars:
            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text and cell.text.strip()
                    )
                    if not row_text:
                        continue
                    if add_block(row_text):
                        break
                if max_chars is not None and total >= max_chars:
                    break

        return "\n\n".join(chunks).strip(), "docx"

    def _normalize_extensions(self, extensions: list[str] | None) -> set[str]:
        if not extensions:
            return set(DEFAULT_TEXT_EXTENSIONS)
        result = set()
        for item in extensions:
            ext = str(item or "").strip().lower()
            if not ext:
                continue
            if not ext.startswith("."):
                ext = "." + ext
            if ext in DEFAULT_TEXT_EXTENSIONS:
                result.add(ext)
        return result

    def _redact_secrets(self, text: str) -> str:
        for pattern in SECRET_PATTERNS:
            text = pattern.sub("[REDACTED_SECRET]", text)
        return text
